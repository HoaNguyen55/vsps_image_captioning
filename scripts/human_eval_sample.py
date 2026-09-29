#!/usr/bin/env python
"""Build the blinded rating sheet for the small human evaluation of caption quality.

    python scripts/human_eval_sample.py --out-dir data/human_eval

50 KTVIC test images (fixed seed) x 3 systems (zero-shot, P2, VSPS seed 42), detailed
mode, verbatim captions. Captions of one image are grouped so the rater opens each
image once; image order and system order inside an image are shuffled; system names
never appear in the sheet. Writes:

  rating_sheet.xlsx  -- give an identical copy to each rater (rater_<name>.xlsx)
  key.json           -- item -> (image_id, system); keep it away from the raters

Scores are entered by people only. Nothing here or in human_eval_score.py
generates a rating.
"""

from __future__ import annotations

import argparse
import json
import os
import random
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SYSTEMS = {
    "zero-shot": "zeroshot-detailed.preds.json",
    "P2": "p2-noverif4090-detailed.preds.json",
    "VSPS": "main_s42-detailed.preds.json",
}
CRITERIA = ("faithfulness", "completeness", "fluency")
RUBRIC = [
    ("Trung thực (faithfulness)", "5 = mọi chi tiết đều đúng với ảnh · 3 = có 1–2 chi tiết sai hoặc bịa · 1 = phần lớn sai/bịa"),
    ("Đầy đủ (completeness)", "5 = nêu được các đối tượng/hành động chính của ảnh · 3 = thiếu một phần quan trọng · 1 = gần như không mô tả nội dung chính"),
    ("Trôi chảy (fluency)", "5 = tiếng Việt tự nhiên, đúng ngữ pháp · 3 = hiểu được nhưng gượng/lặp · 1 = khó hiểu, sai ngữ pháp nặng, lẫn ký tự lạ"),
    ("Quy tắc", "Chấm độc lập từng caption chỉ dựa vào ảnh (không dựa vào caption khác). Không trao đổi với người chấm còn lại."),
]


def default_ktvic() -> Path:
    """$NCS_DATA/datasets/ktvic, else the copy inside the repository (data/datasets/ktvic)."""
    env = Path(os.environ.get("NCS_DATA", Path.home() / "ncs-data")) / "datasets" / "ktvic"
    return env if env.exists() else ROOT / "data" / "datasets" / "ktvic"


def file_names(ktvic: Path) -> dict[str, str]:
    """Exact KTVIC file names when the annotation file is present, else the zero-padded id."""
    ann = ktvic / "test_data.json"
    if ann.exists():
        data = json.loads(ann.read_text(encoding="utf-8"))
        return {str(i.get("id", i.get("image_id"))): i.get("file_name") or i.get("filename") for i in data["images"]}
    return {}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results", default=str(ROOT / "data" / "results"))
    ap.add_argument("--out-dir", default=str(ROOT / "data" / "human_eval"))
    ap.add_argument("--n", type=int, default=50)
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--ktvic", default=str(default_ktvic()),
                    help="KTVIC directory (test_data.json + images/); images are embedded when present")
    args = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.worksheet.datavalidation import DataValidation
    from PIL import Image

    preds = {s: json.loads((Path(args.results) / f).read_text(encoding="utf-8")) for s, f in SYSTEMS.items()}
    common = sorted(set.intersection(*(set(p) for p in preds.values())), key=int)
    rng = random.Random(args.seed)
    images = rng.sample(common, args.n)
    ktvic = Path(args.ktvic)
    names = file_names(ktvic)
    out = Path(args.out_dir)
    thumbs = out / "thumbs"
    missing_images = []

    wb = Workbook()
    ws = wb.active
    ws.title = "rating"
    header = ["item", "image_file", "image", "caption", *CRITERIA, "note"]
    ws.append(header)
    key = {}
    item = 0
    for image_id in images:
        order = list(SYSTEMS)
        rng.shuffle(order)
        name = names.get(image_id, f"{int(image_id):011d}.jpg")
        first_row = ws.max_row + 1
        for system in order:
            item += 1
            cap = preds[system][image_id]
            cap = (cap[0] if isinstance(cap, list) else cap) or ""
            ws.append([item, name, None, cap.strip(), None, None, None, None])
            ws.row_dimensions[ws.max_row].height = 90
            key[str(item)] = {"image_id": image_id, "system": system}
        src = ktvic / "images" / name
        if src.exists():
            thumbs.mkdir(parents=True, exist_ok=True)
            thumb = Image.open(src).convert("RGB")
            thumb.thumbnail((440, 340))
            thumb.save(thumbs / name, quality=85)
            ws.add_image(XlImage(str(thumbs / name)), f"C{first_row}")
        else:
            missing_images.append(name)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:G{item + 1}")
    for col, width in zip("ABCDEFGH", (6, 16, 62, 80, 13, 13, 10, 25)):
        ws.column_dimensions[col].width = width
    group_top = Border(top=Side(style="medium"))
    for r, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if (r - 2) % len(SYSTEMS) == 0:
                c.border = group_top
    for cell in ws[1]:
        cell.font = Font(bold=True)

    guide = wb.create_sheet("huong_dan")
    for name, text in RUBRIC:
        guide.append([name, text])
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 110

    out.mkdir(parents=True, exist_ok=True)
    wb.save(out / "rating_sheet.xlsx")
    (out / "key.json").write_text(json.dumps(
        {"seed": args.seed, "images": images, "systems": SYSTEMS, "items": key},
        ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{args.n} images x {len(SYSTEMS)} systems = {item} items -> {out}/rating_sheet.xlsx (+ key.json)")
    if missing_images:
        print(f"  ! {len(missing_images)} images not found under {ktvic / 'images'} (not embedded): {missing_images[:5]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
