#!/usr/bin/env python
"""Score the human evaluation of caption quality.

    python scripts/human_eval_score.py --key data/human_eval/key.json \
        --sheets data/human_eval/rater_A.xlsx data/human_eval/rater_B.xlsx \
        --out data/results/human_eval.json

Per system and criterion: mean over images of the raters' average, 95% bootstrap CI
(5,000 resamples over images, seed 42). Agreement: quadratic-weighted Cohen's kappa
between the first two raters, per criterion. Tests: two-sided Wilcoxon signed-rank,
VSPS vs every other system, paired by image.
"""

from __future__ import annotations

import argparse
import json
import random
from collections import defaultdict
from pathlib import Path

CRITERIA = ("faithfulness", "completeness", "fluency")


def weighted_kappa(a: list[int], b: list[int], k: int = 5) -> float:
    """Quadratic-weighted Cohen's kappa on 1..k ratings."""
    n = len(a)
    obs = [[0.0] * k for _ in range(k)]
    for x, y in zip(a, b):
        obs[x - 1][y - 1] += 1 / n
    pa = [sum(row) for row in obs]
    pb = [sum(obs[i][j] for i in range(k)) for j in range(k)]
    w = [[((i - j) / (k - 1)) ** 2 for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * obs[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * pa[i] * pb[j] for i in range(k) for j in range(k))
    return 1 - num / den if den else float("nan")


def read_sheet(path: Path) -> dict[str, dict[str, int]]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)["rating"]
    header = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        scores = {c: rec.get(c) for c in CRITERIA}
        missing = [c for c, v in scores.items() if v not in (1, 2, 3, 4, 5)]
        if missing:
            raise SystemExit(f"{path.name}: item {rec['item']} has no valid 1-5 score for {missing}")
        out[str(rec["item"])] = scores
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--key", required=True)
    ap.add_argument("--sheets", nargs="+", required=True)
    ap.add_argument("--out", default="data/results/human_eval.json")
    args = ap.parse_args()

    from scipy.stats import wilcoxon

    key = json.loads(Path(args.key).read_text(encoding="utf-8"))["items"]
    raters = [read_sheet(Path(p)) for p in args.sheets]
    # per (system, criterion): {image_id: mean over raters}
    table: dict[tuple[str, str], dict[str, float]] = defaultdict(dict)
    for item, meta in key.items():
        for c in CRITERIA:
            table[(meta["system"], c)][meta["image_id"]] = sum(r[item][c] for r in raters) / len(raters)

    rng = random.Random(42)
    systems = sorted({m["system"] for m in key.values()})
    images = sorted({m["image_id"] for m in key.values()})
    report: dict = {"n_images": len(images), "n_raters": len(raters), "systems": {}, "kappa": {}, "wilcoxon_vs_VSPS": {}}
    for s in systems:
        for c in CRITERIA:
            v = table[(s, c)]
            boots = sorted(sum(v[rng.choice(images)] for _ in images) / len(images) for _ in range(5000))
            report["systems"].setdefault(s, {})[c] = {
                "mean": round(sum(v.values()) / len(v), 3),
                "ci95": [round(boots[124], 3), round(boots[4874], 3)],
            }
    if len(raters) >= 2:
        for c in CRITERIA:
            a = [raters[0][i][c] for i in key]
            b = [raters[1][i][c] for i in key]
            report["kappa"][c] = round(weighted_kappa(a, b), 3)
    for s in systems:
        if s == "VSPS":
            continue
        for c in CRITERIA:
            x = [table[("VSPS", c)][i] for i in images]
            y = [table[(s, c)][i] for i in images]
            report["wilcoxon_vs_VSPS"].setdefault(s, {})[c] = round(float(wilcoxon(x, y).pvalue), 5)

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
